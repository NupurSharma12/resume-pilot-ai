"""Shared DOCX builders for `app.document_editing` tests.

Not a test file itself (no `test_` prefix, not collected by pytest) --
matches this codebase's precedent of one test module importing a helper
straight from another (e.g. `test_tailoring_suggestions_api.py` importing
`FakeGateway` from `test_tailoring_suggestion_workflow.py`) rather than
duplicating a moderately complex builder three times.
"""

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Pt, RGBColor

from app.models.tailoring_suggestions import (
    SuggestionOperation,
    SuggestionValidationStatus,
    TailoringSuggestion,
)


def build_sample_docx_bytes() -> bytes:
    """A small but formatting-rich resume: distinct fonts/sizes/colors/bold, bullets, a table.

    Deliberately varies formatting per paragraph (not uniform), so tests
    can assert that an *unrelated* paragraph's formatting survives a
    nearby edit untouched, and that an *edited* paragraph's own
    formatting is preserved exactly the way each operation promises to.
    """
    document = Document()

    document.add_heading("SUMMARY", level=2)
    summary_paragraph = document.add_paragraph()
    summary_run = summary_paragraph.add_run("Backend engineer with 5 years of experience.")
    summary_run.font.size = Pt(11)
    summary_run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    document.add_heading("SKILLS", level=2)
    python_paragraph = document.add_paragraph(style="List Bullet")
    python_run = python_paragraph.add_run("Python")
    python_run.bold = True
    python_run.font.size = Pt(12)

    django_paragraph = document.add_paragraph(style="List Bullet")
    django_run = django_paragraph.add_run("Django")
    django_run.italic = True

    document.add_heading("EXPERIENCE", level=2)
    document.add_paragraph(style="List Bullet").add_run(
        "Built internal tools using Python and Django."
    )
    document.add_paragraph(style="List Bullet").add_run("Irrelevant certification from 2005.")

    # A table -- must never be touched by anything in this module, since
    # `DocxStructureMapper` only ever walks top-level body paragraphs.
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Years"

    section = document.sections[0]
    section.header.paragraphs[0].text = "Nupur Sharma — Resume"
    section.footer.paragraphs[0].text = "Page 1"

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def open_docx(content: bytes) -> DocumentObject:
    return Document(BytesIO(content))


def make_suggestion(
    suggestion_id: str,
    target_section_id: str,
    target_item_id: str,
    operation: SuggestionOperation,
    current_text: str | None,
    suggested_text: str,
) -> TailoringSuggestion:
    return TailoringSuggestion(
        suggestion_id=suggestion_id,
        target_section_id=target_section_id,
        target_item_id=target_item_id,
        operation=operation,
        current_text=current_text,
        suggested_text=suggested_text,
        reason="Test reason.",
        evidence_ids=["resume-" + target_item_id],
        evidence_sources=["Resume: test"],
        confidence=90,
        selected_by_default=True,
        validation_status=SuggestionValidationStatus.SUPPORTED_BY_ORIGINAL_RESUME,
        validation_issues=[],
    )
