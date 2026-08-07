"""Unit tests for `DocxParagraphNode`."""

from io import BytesIO

from docx import Document

from app.document_editing.docx_nodes import DocxParagraphNode
from app.document_editing.docx_structure_mapper import DocxStructureMapper
from tests.document_editing_fixtures import build_sample_docx_bytes, open_docx


def _node_for(document, text: str) -> DocxParagraphNode:
    _, nodes = DocxStructureMapper().map(document)
    for node in nodes.values():
        if node.get_text() == text:
            return node
    raise AssertionError(f"No node found with text {text!r}")


def _roundtrip(document) -> Document:
    """Save `document` and reopen it, the same way the real editor does."""
    buffer = BytesIO()
    document.save(buffer)
    return open_docx(buffer.getvalue())


def test_set_text_replaces_a_single_run_paragraphs_text_and_keeps_its_formatting() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Python")

    node.set_text("Python, TypeScript")
    reopened = _roundtrip(document)

    paragraph = next(p for p in reopened.paragraphs if p.text == "Python, TypeScript")
    assert len(paragraph.runs) == 1
    assert paragraph.runs[0].bold is True


def test_append_text_adds_a_new_run_and_never_touches_the_existing_one() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Python")

    node.append_text(", TypeScript")
    reopened = _roundtrip(document)

    paragraph = next(p for p in reopened.paragraphs if p.text == "Python, TypeScript")
    assert len(paragraph.runs) == 2
    assert paragraph.runs[0].text == "Python"
    assert paragraph.runs[0].bold is True
    # The new run clones the *previous last* run's formatting.
    assert paragraph.runs[1].text == ", TypeScript"
    assert paragraph.runs[1].bold is True


def test_append_text_with_empty_addition_is_a_no_op() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Python")

    node.append_text("")

    assert node.get_text() == "Python"
    assert len(node._paragraph.runs) == 1


def test_insert_after_clones_the_anchors_style_and_formatting() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Python")

    new_node = node.insert_after("Node.js")
    reopened = _roundtrip(document)

    paragraph = next(p for p in reopened.paragraphs if p.text == "Node.js")
    assert paragraph.style.name == "List Bullet"
    assert paragraph.runs[0].bold is True  # cloned from the "Python" run
    assert new_node.get_text() == "Node.js"
    # It must land immediately after "Python", before "Django".
    texts = [p.text for p in reopened.paragraphs]
    assert texts.index("Python") < texts.index("Node.js") < texts.index("Django")


def test_insert_before_places_the_new_node_immediately_before_the_anchor() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Django")

    node.insert_before("Node.js")
    reopened = _roundtrip(document)

    texts = [p.text for p in reopened.paragraphs]
    assert texts.index("Python") < texts.index("Node.js") < texts.index("Django")


def test_remove_deletes_the_paragraph_entirely() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Irrelevant certification from 2005.")

    node.remove()
    reopened = _roundtrip(document)

    texts = [p.text for p in reopened.paragraphs]
    assert "Irrelevant certification from 2005." not in texts
    # Its sibling must still be present, untouched.
    assert "Built internal tools using Python and Django." in texts


def test_set_text_on_a_paragraph_with_no_runs_adds_one() -> None:
    document = Document()
    empty_paragraph = document.add_paragraph()
    node = DocxParagraphNode("item-0", empty_paragraph)

    node.set_text("New text")

    assert node.get_text() == "New text"


def test_get_text_reflects_live_paragraph_state() -> None:
    document = open_docx(build_sample_docx_bytes())
    node = _node_for(document, "Django")

    assert node.get_text() == "Django"
    node.set_text("Django, Flask")
    assert node.get_text() == "Django, Flask"
