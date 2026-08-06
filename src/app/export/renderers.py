"""Per-format renderers: turn a `StructuredResume` into downloadable file bytes.

Every renderer here is a pure function of `StructuredResume` — none of them
read the original uploaded file, because none of them can: see
`docs/features/interactive-tailored-resume.md` for why no original file
bytes are retained anywhere past upload. Each renderer produces the
*content* of the final resume in its target format; `ExportService` (not
this module) decides filenames, content types, and how to honestly label
each result's fidelity to the original.
"""

from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate

from app.models.resume_structure import StructuredResume


class TxtResumeRenderer:
    """Renders a `StructuredResume` as clean plain text."""

    def render(self, structured_resume: StructuredResume) -> bytes:
        return structured_resume.to_text().encode("utf-8")


class MarkdownResumeRenderer:
    """Renders a `StructuredResume` as clean Markdown (`##` headings, `-` bullets)."""

    def render(self, structured_resume: StructuredResume) -> bytes:
        lines: list[str] = []
        for section in structured_resume.sections:
            if section.heading:
                lines.append(f"## {section.heading}")
            for item in section.items:
                lines.append(f"- {item.text}")
            lines.append("")
        text = "\n".join(lines).rstrip() + "\n"
        return text.encode("utf-8")


class DocxResumeRenderer:
    """Renders a `StructuredResume` as a fresh DOCX document in a clean default style.

    This is always a newly generated document — the original DOCX's
    fonts, colors, and layout are never available to reproduce (only its
    extracted paragraph text survived upload). `ExportService` labels
    this `FormatFidelity.REGENERATED_TEMPLATE` accordingly, never
    `APPROXIMATE_STYLE`.
    """

    def render(self, structured_resume: StructuredResume) -> bytes:
        document = Document()
        for section in structured_resume.sections:
            if section.heading:
                document.add_heading(section.heading, level=2)
            for item in section.items:
                document.add_paragraph(item.text, style="List Bullet")
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()


class PdfResumeRenderer:
    """Renders a `StructuredResume` as a simple, readable, ATS-friendly PDF.

    Always a freshly generated layout, never a reproduction of any
    original PDF's design — see `docs/features/interactive-tailored-resume.md`
    for why this application cannot promise source-layout preservation
    for PDF. `ExportService` labels every PDF export
    `FormatFidelity.REGENERATED_TEMPLATE`, unconditionally.
    """

    def render(self, structured_resume: StructuredResume) -> bytes:
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=LETTER,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
            leftMargin=0.75 * inch,
            rightMargin=0.75 * inch,
        )
        styles = getSampleStyleSheet()
        heading_style = ParagraphStyle(
            "ResumeSectionHeading",
            parent=styles["Heading2"],
            spaceBefore=12,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "ResumeBody",
            parent=styles["Normal"],
            spaceAfter=4,
        )

        story = []
        for section in structured_resume.sections:
            if section.heading:
                story.append(Paragraph(_escape(section.heading), heading_style))
            for item in section.items:
                story.append(Paragraph(f"&bull; {_escape(item.text)}", body_style))
        if not story:
            story.append(Paragraph("(empty resume)", body_style))

        doc.build(story)
        return buffer.getvalue()


def _escape(text: str) -> str:
    """Escape XML-significant characters before embedding `text` in a reportlab `Paragraph`.

    `Paragraph` interprets its input as a small XML-like markup language,
    so raw resume text containing `&`, `<`, or `>` must be escaped first
    or it would be misread as markup rather than rendered literally.
    """
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
